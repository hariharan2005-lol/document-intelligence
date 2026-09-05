"""Script to generate realistic test fixtures (PDF, DOCX, PNG) for invoices and resumes."""
import os
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

FIXTURES_DIR = Path(__file__).parent


def create_invoice_pdf(output_path: Path):
    """Generate a realistic PDF invoice."""
    c = canvas.Canvas(str(output_path), pagesize=letter)
    
    # Header / Company Info
    c.setFont("Helvetica-Bold", 20)
    c.drawString(50, 750, "Acme Solutions Corp")
    
    c.setFont("Helvetica", 10)
    c.drawString(50, 735, "100 Enterprise Way, Suite 400")
    c.drawString(50, 722, "Tech City, CA 94016")
    c.drawString(50, 709, "billing@acmesolutions.com")
    
    # Invoice details
    c.setFont("Helvetica-Bold", 14)
    c.drawString(400, 750, "INVOICE")
    
    c.setFont("Helvetica", 10)
    c.drawString(400, 735, "Invoice Number: INV-2024-8842")
    c.drawString(400, 722, "Date: 2024-03-15")
    c.drawString(400, 709, "Due Date: 2024-04-15")
    
    # Bill To
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, 660, "Bill To:")
    c.setFont("Helvetica", 10)
    c.drawString(50, 645, "Globex International")
    c.drawString(50, 632, "Attn: Accounts Payable")
    c.drawString(50, 619, "500 Global Plaza, New York, NY 10001")
    
    # Line items table header
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 560, "Description")
    c.drawString(320, 560, "Quantity")
    c.drawString(400, 560, "Unit Price")
    c.drawString(480, 560, "Total")
    c.line(50, 552, 540, 552)
    
    # Line items
    c.setFont("Helvetica", 10)
    c.drawString(50, 535, "Cloud Architecture Consulting")
    c.drawString(340, 535, "10")
    c.drawString(400, 535, "$150.00")
    c.drawString(480, 535, "$1500.00")
    
    c.drawString(50, 515, "API Ingestion Pipeline Development")
    c.drawString(340, 515, "1")
    c.drawString(400, 515, "$3000.00")
    c.drawString(480, 515, "$3000.00")
    
    # Totals
    c.line(50, 495, 540, 495)
    c.drawString(400, 475, "Subtotal:")
    c.drawString(480, 475, "$4500.00")
    c.drawString(400, 455, "Tax (0%):")
    c.drawString(480, 455, "$0.00")
    
    c.setFont("Helvetica-Bold", 12)
    c.drawString(400, 430, "Total Amount Due:")
    c.drawString(480, 430, "$4500.00")
    
    # Footer / Boilerplate to test cleaner
    c.setFont("Helvetica", 8)
    c.drawString(250, 50, "Page 1 of 1")
    c.drawString(50, 35, "Confidential - All rights reserved. Thank you for your business.")
    
    c.save()


def create_resume_pdf(output_path: Path):
    """Generate a realistic PDF candidate resume."""
    c = canvas.Canvas(str(output_path), pagesize=letter)
    
    # Header
    c.setFont("Helvetica-Bold", 22)
    c.drawString(50, 750, "Alex Mercer")
    
    c.setFont("Helvetica", 10)
    c.drawString(50, 732, "alex.mercer@devmail.com | (555) 234-5678 | San Francisco, CA | github.com/alexmercer")
    c.line(50, 722, 550, 722)
    
    # Summary
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 700, "PROFESSIONAL SUMMARY")
    c.setFont("Helvetica", 10)
    c.drawString(50, 685, "Senior Backend Engineer with 6+ years of experience designing scalable data pipelines,")
    c.drawString(50, 672, "microservices, and RESTful APIs using Python, FastAPI, Flask, and PostgreSQL.")
    
    # Skills
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 640, "SKILLS")
    c.setFont("Helvetica", 10)
    c.drawString(50, 625, "Languages: Python, SQL, TypeScript, Bash")
    c.drawString(50, 612, "Frameworks & Libraries: FastAPI, Flask, Django, SQLAlchemy, PyTorch")
    c.drawString(50, 599, "Tools & Cloud: Docker, Kubernetes, AWS, Git, Redis, PostgreSQL")
    
    # Experience
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 565, "WORK EXPERIENCE")
    
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 550, "Senior Backend Engineer | DataCorp Solutions")
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(430, 550, "2021 - Present")
    c.setFont("Helvetica", 9)
    c.drawString(60, 535, "• Built asynchronous document intelligence microservices processing 50k documents daily.")
    c.drawString(60, 522, "• Developed high-performance REST APIs with FastAPI and Flask.")
    
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 495, "Software Engineer | CloudScale Inc")
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(430, 495, "2018 - 2021")
    c.setFont("Helvetica", 9)
    c.drawString(60, 480, "• Developed scalable web services and background queue workers in Python and PostgreSQL.")
    
    # Education
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 445, "EDUCATION")
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 430, "Bachelor of Science in Computer Science | University of California, Berkeley")
    c.setFont("Helvetica-Oblique", 9)
    c.drawString(450, 430, "2018")
    
    # Footer to test cleaner
    c.setFont("Helvetica", 8)
    c.drawString(270, 30, "Page 1 of 1")
    
    c.save()


def create_invoice_docx(output_path: Path):
    """Generate a realistic DOCX invoice."""
    doc = docx.Document()
    
    doc.add_heading("Apex Consulting Ltd", level=1)
    doc.add_paragraph("Invoice Number: INV-DOCX-9901")
    doc.add_paragraph("Date: 2024-02-20")
    doc.add_paragraph("Bill To: Acme Global Enterprises")
    
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Item"
    hdr_cells[1].text = "Quantity"
    hdr_cells[2].text = "Amount"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "Software Architecture Review"
    row_cells[1].text = "1"
    row_cells[2].text = "$2800.00"
    
    doc.add_paragraph("Total Amount Due: $2800.00 USD")
    doc.add_paragraph("Page 1 of 1")
    doc.save(str(output_path))


def create_resume_docx(output_path: Path):
    """Generate a realistic DOCX resume."""
    doc = docx.Document()
    
    doc.add_heading("Sarah Connor", level=1)
    doc.add_paragraph("sarah.connor@cyberdyne.org | 555-432-8765")
    
    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, Flask, Docker, Kubernetes, Linux, AWS, SQL")
    
    doc.add_heading("Work Experience", level=2)
    doc.add_paragraph("DevOps Engineer at Cyberdyne Systems (2020 - Present)")
    doc.add_paragraph("Software Engineer at TechCore (2017 - 2020)")
    
    doc.add_heading("Education", level=2)
    doc.add_paragraph("Master of Science in Computer Engineering, Stanford University, 2017")
    
    doc.save(str(output_path))


def create_invoice_png(output_path: Path):
    """Generate a PNG image invoice fixture."""
    img = Image.new("RGB", (600, 400), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    
    # Draw simple text lines
    draw.text((30, 30), "INVOICE - Swift Courier Services", fill=(0, 0, 0))
    draw.text((30, 60), "Invoice Number: INV-IMG-7731", fill=(0, 0, 0))
    draw.text((30, 90), "Date: 2024-04-10", fill=(0, 0, 0))
    draw.text((30, 120), "Bill To: Delta Logistics", fill=(0, 0, 0))
    draw.text((30, 150), "Total Due: $650.00 USD", fill=(0, 0, 0))
    draw.text((30, 350), "Page 1 of 1", fill=(100, 100, 100))
    
    img.save(str(output_path), "PNG")


def create_invoice_variant_pdf(output_path: Path):
    """Generate a realistic PDF invoice with alternative field label wording (Ref No, DD-Mon-YYYY, Grand Total)."""
    c = canvas.Canvas(str(output_path), pagesize=letter)
    
    # Header / Vendor
    c.setFont("Helvetica-Bold", 18)
    c.drawString(50, 750, "Northwind Traders")
    c.setFont("Helvetica", 10)
    c.drawString(50, 735, "456 Commerce Boulevard, Dock 12")
    c.drawString(50, 722, "accounts@northwindtraders.com")
    
    # Document Reference & Date with alternative wording
    c.setFont("Helvetica-Bold", 12)
    c.drawString(380, 750, "TAX INVOICE")
    c.setFont("Helvetica", 10)
    c.drawString(380, 735, "Ref No: NM-77821")
    c.drawString(380, 722, "Date: 02-Sep-2026")
    c.drawString(380, 709, "Payment Terms: Net 30")
    
    # Customer Info
    c.setFont("Helvetica-Bold", 11)
    c.drawString(50, 660, "Bill To:")
    c.setFont("Helvetica", 10)
    c.drawString(50, 645, "Contoso Ltd")
    c.drawString(50, 632, "Attn: Procurement Division")
    c.drawString(50, 619, "789 Enterprise Blvd, Suite 200")
    
    # Line items table
    c.setFont("Helvetica-Bold", 10)
    c.drawString(50, 560, "Item")
    c.drawString(320, 560, "Qty")
    c.drawString(400, 560, "Rate")
    c.drawString(480, 560, "Total")
    c.line(50, 552, 540, 552)
    
    c.setFont("Helvetica", 10)
    c.drawString(50, 535, "Enterprise Cloud Storage Appliance")
    c.drawString(330, 535, "2")
    c.drawString(400, 535, "5,000.00")
    c.drawString(480, 535, "10,000.00")
    
    c.drawString(50, 515, "Annual Infrastructure Support")
    c.drawString(330, 515, "1")
    c.drawString(400, 515, "2,340.50")
    c.drawString(480, 515, "2,340.50")
    
    # Summary amounts
    c.line(50, 495, 540, 495)
    c.drawString(380, 475, "Subtotal:")
    c.drawString(480, 475, "12,340.50")
    
    c.setFont("Helvetica-Bold", 12)
    c.drawString(350, 445, "Grand Total:")
    c.drawString(470, 445, "12,340.50 USD")
    
    c.setFont("Helvetica", 8)
    c.drawString(250, 40, "Page 1 of 1")
    c.save()


def main():
    FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    create_invoice_pdf(FIXTURES_DIR / "sample_invoice.pdf")
    create_resume_pdf(FIXTURES_DIR / "sample_resume.pdf")
    create_invoice_docx(FIXTURES_DIR / "sample_invoice.docx")
    create_resume_docx(FIXTURES_DIR / "sample_resume.docx")
    create_invoice_png(FIXTURES_DIR / "sample_invoice.png")
    create_invoice_variant_pdf(FIXTURES_DIR / "sample_invoice_variant.pdf")
    print("Successfully created test fixtures in:", FIXTURES_DIR)


if __name__ == "__main__":
    main()
